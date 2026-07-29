import csv
import io
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path

from flask import Blueprint, current_app, g, jsonify, request, send_from_directory
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from app.auth import can_manage_accounts, login_required


talk_library_bp = Blueprint("talk_library", __name__, url_prefix="/api/talk-library")

LEARNING_CALL_TITLES = ["首通电话", "第二通电话", "第三通电话", "第四通电话", "第五通电话"]
LEARNING_SECTION_KEYS = ("probe", "output", "concept")
DEFAULT_TALK_CATEGORY_NAMES = ["参课", "催课", "答疑", "续费", "学情", "挽单", "转介绍", "素材库"]
RETAINED_TALK_CATEGORY_NAMES = []
TALK_CATEGORY_ALIASES = {"其他": "答疑"}
DEFAULT_TALK_CATEGORY_DESCRIPTIONS = {
    "参课": "参课沟通与到课提醒",
    "催课": "完课提醒与补课沟通",
    "答疑": "家长常见问题回复",
    "续费": "续费沟通与留言推荐",
    "学情": "学情电话与学习反馈",
    "挽单": "退费挽留与风险沟通",
    "转介绍": "转介绍邀约与报名沟通",
    "素材库": "截图、好评与图片素材",
}
TALK_TRACK_CATEGORIES = set(DEFAULT_TALK_CATEGORY_NAMES) | set(RETAINED_TALK_CATEGORY_NAMES)
TALK_MATERIAL_CATEGORIES = TALK_TRACK_CATEGORIES
TALK_MATERIAL_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
SHARED_TALK_CATEGORIES = {"答疑"}
RENEWAL_TALK_TYPES = {"问答话术", "留言推荐"}
TALK_TRACK_STATUSES = {"启用", "停用"}
TALK_LIBRARY_SCHEMA_VERSION = 2


def talk_library_file():
    return current_app.config["TALK_LIBRARY_FILE"]


def talk_material_dir():
    return Path(current_app.config["TALK_MATERIAL_DIR"])


def blank_talk_library_data():
    return {
        "schema_version": TALK_LIBRARY_SCHEMA_VERSION,
        "categories": default_talk_categories(),
        "learning_calls": {},
        "materials": [],
        "tracks": [],
        "personal_tracks": {},
    }


def default_talk_categories():
    categories = []
    for index, name in enumerate([*DEFAULT_TALK_CATEGORY_NAMES, *RETAINED_TALK_CATEGORY_NAMES], start=1):
        categories.append({
            "id": uuid.uuid5(uuid.NAMESPACE_URL, f"pu6-talk-category:{name}").hex,
            "name": name,
            "sort": index * 10,
            "status": "启用",
            "note": DEFAULT_TALK_CATEGORY_DESCRIPTIONS.get(name, ""),
        })
    return categories


def normalize_track_status(value):
    status = str(value or "").strip()
    if status in {"停用", "禁用", "disabled", "disable", "0", "false", "False"}:
        return "停用"
    return "启用"


def parse_int(value, default=10):
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def first_text(*values):
    for value in values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def canonical_talk_category_name(value):
    category = str(value or "").strip()
    return TALK_CATEGORY_ALIASES.get(category, category)


def normalize_talk_category_item(category, fallback_sort=999):
    if isinstance(category, dict):
        name = canonical_talk_category_name(category.get("name") or category.get("分类") or category.get("category"))
        sort = parse_int(category.get("sort") or category.get("排序"), fallback_sort)
        status = normalize_track_status(category.get("status") or category.get("状态"))
        note = str(category.get("note") or category.get("备注") or DEFAULT_TALK_CATEGORY_DESCRIPTIONS.get(name, "")).strip()
        category_id = str(category.get("id") or "").strip()
    else:
        name = canonical_talk_category_name(category)
        sort = fallback_sort
        status = "启用"
        note = DEFAULT_TALK_CATEGORY_DESCRIPTIONS.get(name, "")
        category_id = ""
    if not name:
        return None
    return {
        "id": category_id or uuid.uuid5(uuid.NAMESPACE_URL, f"pu6-talk-category:{name}").hex,
        "name": name[:40],
        "sort": sort,
        "status": status,
        "note": note[:240],
    }


def normalize_talk_categories(data):
    categories_by_name = {}
    for category in default_talk_categories():
        categories_by_name[category["name"]] = category

    raw_categories = data.get("categories") if isinstance(data.get("categories"), list) else []
    for index, raw_category in enumerate(raw_categories, start=100):
        raw_name = raw_category.get("name") if isinstance(raw_category, dict) else raw_category
        was_alias = str(raw_name or "").strip() in TALK_CATEGORY_ALIASES
        category = normalize_talk_category_item(raw_category, fallback_sort=index)
        if category:
            if was_alias and category["name"] in categories_by_name:
                continue
            categories_by_name[category["name"]] = category

    for track in data.get("tracks", []):
        if not isinstance(track, dict):
            continue
        category_name = canonical_talk_category_name(first_text(track.get("category"), track.get("分类")))
        if category_name and category_name not in categories_by_name:
            categories_by_name[category_name] = normalize_talk_category_item(category_name, fallback_sort=900 + len(categories_by_name))

    ordered = sorted(categories_by_name.values(), key=lambda item: (item.get("sort", 999), item.get("name", "")))
    return ordered


def talk_category_names(data, include_disabled=True):
    names = []
    for category in data.get("categories", []):
        item = normalize_talk_category_item(category)
        if not item:
            continue
        if include_disabled or item.get("status") != "停用":
            names.append(item["name"])
    return names


def normalize_talk_library_data(data):
    if not isinstance(data, dict):
        data = blank_talk_library_data()

    data.setdefault("learning_calls", {})
    data.setdefault("materials", [])
    data.setdefault("tracks", [])
    data.setdefault("personal_tracks", {})
    data["categories"] = normalize_talk_categories(data)

    category_names = set(talk_category_names(data))
    normalized_tracks = []
    for track in data.get("tracks", []):
        item = normalize_talk_track(track, category_names=category_names)
        if item is not None:
            normalized_tracks.append(item)
    data["tracks"] = normalized_tracks
    data["schema_version"] = TALK_LIBRARY_SCHEMA_VERSION
    return data


def load_talk_library():
    path = talk_library_file()
    if not path.exists():
        return blank_talk_library_data()
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)
    original = json.dumps(data, ensure_ascii=False, sort_keys=True)
    normalized = normalize_talk_library_data(data)
    if json.dumps(normalized, ensure_ascii=False, sort_keys=True) != original:
        backup_path = path.with_name(f"{path.name}.bak-unified-v2")
        if not backup_path.exists():
            backup_path.write_bytes(path.read_bytes())
        save_talk_library(normalized)
    return normalized


def save_talk_library(data):
    path = talk_library_file()
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False, indent=2)


def normalize_learning_call(payload):
    title = str(payload.get("title") or "").strip()
    if title not in LEARNING_CALL_TITLES:
        return None

    return {
        "title": title,
        "probe": str(payload.get("probe") or "").strip(),
        "output": str(payload.get("output") or "").strip(),
        "concept": str(payload.get("concept") or "").strip(),
        "updated_by": g.user.get("username", ""),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def public_learning_call(call):
    return {
        "title": call.get("title", ""),
        "probe": call.get("probe", ""),
        "output": call.get("output", ""),
        "concept": call.get("concept", ""),
        "updated_by": call.get("updated_by", ""),
        "updated_at": call.get("updated_at", ""),
    }


def normalize_talk_category(value, category_names=None, default="续费"):
    category = canonical_talk_category_name(value)
    known_categories = set(category_names or TALK_TRACK_CATEGORIES)
    return category if category in known_categories else default


def normalize_renewal_talk_type(value):
    talk_type = str(value or "").strip()
    return talk_type if talk_type in RENEWAL_TALK_TYPES else "问答话术"


def normalize_talk_track(track, category_names=None, require_fields=False):
    if not isinstance(track, dict):
        return None
    text = first_text(
        track.get("content"),
        track.get("话术内容"),
        track.get("text"),
        track.get("标准话术"),
        track.get("talktrack"),
        track.get("answer"),
    )
    if not text:
        return None
    known_categories = set(category_names or TALK_TRACK_CATEGORIES)
    raw_category = canonical_talk_category_name(first_text(track.get("category"), track.get("分类")))
    category = normalize_talk_category(raw_category, category_names=known_categories)
    if require_fields and (not raw_category or raw_category not in known_categories):
        return None

    keyword = first_text(
        track.get("keyword"),
        track.get("关键词"),
        track.get("keywords"),
        track.get("scene"),
        track.get("场景"),
        track.get("example"),
        track.get("问题示例"),
    )
    if require_fields and not keyword:
        return None
    if not keyword:
        keyword = text.replace("\r", "\n").split("\n", 1)[0][:24] or "新增话术"

    track_id = str(track.get("id") or "").strip() or uuid.uuid4().hex
    priority = parse_int(track.get("sort") or track.get("排序") or track.get("priority") or track.get("优先级"), 10)
    status = normalize_track_status(track.get("status") or track.get("状态"))
    note = str(track.get("note") or track.get("备注") or "").strip()
    scene = first_text(track.get("scene"), track.get("场景"), keyword)
    example = first_text(track.get("example"), track.get("问题示例"))
    output = {
        "id": track_id,
        "category": category,
        "type": normalize_renewal_talk_type(track.get("type") or track.get("类型")) if category == "续费" else "",
        "keyword": keyword[:200],
        "keywords": keyword[:200],
        "scene": scene[:120],
        "example": example[:180],
        "content": text[:5000],
        "text": text[:5000],
        "sort": priority,
        "priority": priority,
        "status": status,
        "note": note[:500],
        "created_by": str(track.get("created_by") or track.get("创建人") or "").strip(),
        "created_at": str(track.get("created_at") or "").strip(),
        "updated_at": str(track.get("updated_at") or "").strip(),
    }
    return output


def public_talk_track(track):
    item = normalize_talk_track(track)
    if item is None:
        return None
    output = {
        **item,
        "can_delete": can_delete_talk_track_item(item),
    }
    if not can_manage_accounts():
        output.pop("note", None)
    return output


def can_create_talk_track_item(track):
    if can_manage_accounts():
        return True
    item = normalize_talk_track(track)
    return bool(item and item.get("category") in SHARED_TALK_CATEGORIES)


def can_delete_talk_track_item(track):
    if can_manage_accounts():
        return True
    item = normalize_talk_track(track)
    if item is None or item.get("category") not in SHARED_TALK_CATEGORIES:
        return False
    return item.get("created_by") == str(g.user.get("username") or "")


def normalized_talk_tracks(data):
    tracks = []
    category_names = set(talk_category_names(data))
    for track in data.get("tracks", []):
        item = normalize_talk_track(track, category_names=category_names)
        if item is not None:
            tracks.append(item)
    tracks.sort(key=lambda item: (item.get("sort") or item.get("priority") or 10, item.get("created_at") or "", item.get("keyword") or ""))
    return tracks


def normalize_personal_talk_track(track):
    if not isinstance(track, dict):
        return None
    text = str(
        track.get("text")
        or track.get("标准话术")
        or track.get("talktrack")
        or track.get("answer")
        or ""
    ).strip()
    if not text:
        return None
    track_id = str(track.get("id") or "").strip() or uuid.uuid4().hex
    try:
        priority = int(track.get("priority") or track.get("优先级") or 10)
    except (TypeError, ValueError):
        priority = 10
    raw_category = canonical_talk_category_name(track.get("category") or track.get("分类"))
    category = raw_category if raw_category in TALK_TRACK_CATEGORIES else ""
    output = {
        "id": track_id,
        "category": category,
        "type": "个人常用",
        "scene": str(track.get("scene") or track.get("场景") or track.get("title") or track.get("问题示例") or "").strip()[:120],
        "keywords": str(track.get("keywords") or track.get("关键词") or "").strip()[:200],
        "example": str(track.get("example") or track.get("问题示例") or "").strip()[:180],
        "text": text[:5000],
        "priority": priority,
        "created_by": str(track.get("created_by") or "").strip(),
        "created_at": str(track.get("created_at") or "").strip(),
        "updated_at": str(track.get("updated_at") or "").strip(),
    }
    if not output["scene"]:
        output["scene"] = output["keywords"] or text[:24] or "常用话术"
    return output


def public_personal_talk_track(track):
    item = normalize_personal_talk_track(track)
    if item is None:
        return None
    return {
        **item,
        "personal": True,
        "can_delete": True,
    }


def personal_talk_bucket(data, username):
    buckets = data.setdefault("personal_tracks", {})
    if not isinstance(buckets, dict):
        data["personal_tracks"] = {}
        buckets = data["personal_tracks"]
    username = str(username or "").strip()
    buckets.setdefault(username, [])
    if not isinstance(buckets[username], list):
        buckets[username] = []
    return buckets[username]


def normalized_personal_talk_tracks(data, username):
    tracks = []
    for track in personal_talk_bucket(data, username):
        item = normalize_personal_talk_track(track)
        if item is not None:
            tracks.append(item)
    tracks.sort(key=lambda item: (item.get("created_at") or "", item.get("priority") or 0), reverse=True)
    return tracks


def normalize_material_category(value):
    category = str(value or "").strip()
    return category if category in TALK_MATERIAL_CATEGORIES else "素材库"


def normalize_material(material):
    if not isinstance(material, dict):
        return None
    material_id = str(material.get("id") or "").strip()
    filename = str(material.get("filename") or "").strip()
    if not material_id or not filename:
        return None
    title = str(material.get("title") or "").strip() or str(material.get("original_filename") or "").strip() or "素材"
    category = normalize_material_category(material.get("category"))
    return {
        "id": material_id,
        "title": title[:80],
        "category": category,
        "keywords": str(material.get("keywords") or "").strip()[:160],
        "note": str(material.get("note") or "").strip()[:240],
        "filename": filename,
        "original_filename": str(material.get("original_filename") or "").strip(),
        "mime_type": str(material.get("mime_type") or "").strip(),
        "size": int(material.get("size") or 0),
        "uploaded_by": str(material.get("uploaded_by") or "").strip(),
        "uploaded_at": str(material.get("uploaded_at") or "").strip(),
    }


def public_material(material):
    item = normalize_material(material)
    if item is None:
        return None
    return {
        **item,
        "can_delete": can_manage_accounts(),
        "url": f"/api/talk-library/materials/{item['id']}/file",
    }


def normalized_materials(data):
    materials = []
    for material in data.get("materials", []):
        item = normalize_material(material)
        if item is not None:
            materials.append(item)
    materials.sort(key=lambda item: item.get("uploaded_at") or "", reverse=True)
    return materials


def public_talk_category(category):
    item = normalize_talk_category_item(category)
    if item is None:
        return None
    output = {
        "id": item["id"],
        "name": item["name"],
        "sort": item["sort"],
        "status": item["status"],
    }
    if can_manage_accounts():
        output["note"] = item.get("note", "")
    return output


def desktop_talk_track(track):
    item = normalize_talk_track(track)
    if item is None or item.get("status") == "停用":
        return None
    return {
        "id": item["id"],
        "category": item["category"],
        "keyword": item["keyword"],
        "content": item["content"],
        "sort": item["sort"],
        "status": item["status"],
    }


def desktop_talk_category(category):
    item = normalize_talk_category_item(category)
    if item is None or item.get("status") == "停用":
        return None
    return {
        "id": item["id"],
        "name": item["name"],
        "sort": item["sort"],
        "status": item["status"],
    }


def extract_talk_track_import_row(raw_track, row_number, category_names, now):
    category = canonical_talk_category_name(first_text(raw_track.get("category"), raw_track.get("分类")))
    keyword = first_text(raw_track.get("keyword"), raw_track.get("关键词"), raw_track.get("keywords"))
    content = first_text(raw_track.get("content"), raw_track.get("话术内容"), raw_track.get("text"), raw_track.get("标准话术"))

    errors = []
    if not category:
        errors.append("分类必填")
    elif category not in category_names:
        errors.append("分类不存在，请先在工作集里创建或修正分类名")
    if not keyword:
        errors.append("关键词必填")
    if not content:
        errors.append("话术内容必填")

    if errors:
        return None, {
            "row": row_number,
            "category": category,
            "keyword": keyword,
            "errors": errors,
        }

    track = normalize_talk_track({
        "category": category,
        "keyword": keyword,
        "content": content,
        "sort": raw_track.get("sort") or raw_track.get("排序"),
        "status": raw_track.get("status") or raw_track.get("状态"),
        "note": raw_track.get("note") or raw_track.get("备注"),
    }, category_names=category_names, require_fields=True)
    if track is None:
        return None, {
            "row": row_number,
            "category": category,
            "keyword": keyword,
            "errors": ["没有读取到可导入的话术"],
        }
    track["id"] = uuid.uuid4().hex
    track["created_by"] = str(g.user.get("username") or "")
    track["created_at"] = now
    track["updated_at"] = now
    return track, None


def import_talk_track_rows(data, raw_tracks, replace_category="", should_replace_category=False):
    category_names = set(talk_category_names(data))
    now = datetime.now(timezone.utc).isoformat()
    imported = []
    invalid_rows = []
    for row_number, raw_track in enumerate(raw_tracks, start=2):
        if not isinstance(raw_track, dict):
            invalid_rows.append({"row": row_number, "errors": ["行数据格式错误"]})
            continue
        if not any(str(value or "").strip() for value in raw_track.values()):
            continue
        track, error = extract_talk_track_import_row(raw_track, row_number, category_names, now)
        if error:
            invalid_rows.append(error)
        elif track:
            imported.append(track)

    if invalid_rows:
        return imported, invalid_rows

    if not imported:
        return imported, [{"row": 0, "errors": ["没有读取到可导入的话术。"]}]

    tracks = normalized_talk_tracks(data)
    if should_replace_category and replace_category in category_names:
        tracks = [track for track in tracks if track.get("category") != replace_category]
    data["tracks"] = [*imported, *tracks]
    return imported, []


def parse_uploaded_talk_file(uploaded_file):
    raw_filename = str(uploaded_file.filename or "").strip()
    extension = Path(raw_filename).suffix.lower()
    if extension == ".xlsx":
        uploaded_file.stream.seek(0)
        workbook = load_workbook(uploaded_file.stream, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        parsed_rows = []
        for row in rows[1:]:
            parsed_rows.append({
                headers[index]: value
                for index, value in enumerate(row)
                if index < len(headers) and headers[index]
            })
        return parsed_rows
    if extension == ".csv":
        uploaded_file.stream.seek(0)
        raw_bytes = uploaded_file.read()
        try:
            text = raw_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw_bytes.decode("gb18030", errors="replace")
        return list(csv.DictReader(io.StringIO(text)))
    raise ValueError("目前仅支持 xlsx 或 csv 文件。")


@talk_library_bp.get("/categories")
@login_required
def get_talk_categories():
    data = load_talk_library()
    categories = []
    for category in data.get("categories", []):
        item = normalize_talk_category_item(category)
        if item is None:
            continue
        if can_manage_accounts() or item.get("status") != "停用":
            categories.append(public_talk_category(item))
    return jsonify({"categories": [category for category in categories if category is not None]})


@talk_library_bp.put("/categories")
@login_required
def save_talk_categories():
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以维护话术分类。"}), 403

    payload = request.get_json(silent=True) or {}
    raw_categories = payload.get("categories") if isinstance(payload.get("categories"), list) else []
    if not raw_categories:
        return jsonify({"error": "请至少保留一个话术分类。"}), 400

    categories = []
    seen_names = set()
    for index, raw_category in enumerate(raw_categories, start=1):
        category = normalize_talk_category_item(raw_category, fallback_sort=index * 10)
        if category is None:
            continue
        if category["name"] in seen_names:
            return jsonify({"error": f"分类重复：{category['name']}"}), 400
        seen_names.add(category["name"])
        categories.append(category)
    if not categories:
        return jsonify({"error": "请至少保留一个话术分类。"}), 400

    data = load_talk_library()
    data["categories"] = categories
    save_talk_library(data)
    return jsonify({"categories": [public_talk_category(category) for category in categories]})


@talk_library_bp.get("/desktop")
@login_required
def get_desktop_talk_library():
    data = load_talk_library()
    categories = [
        desktop_talk_category(category)
        for category in data.get("categories", [])
    ]
    tracks = [
        desktop_talk_track(track)
        for track in normalized_talk_tracks(data)
    ]
    return jsonify({
        "categories": [category for category in categories if category is not None],
        "tracks": [track for track in tracks if track is not None],
    })


@talk_library_bp.get("/learning-calls")
@login_required
def get_learning_calls():
    data = load_talk_library()
    calls = {
        title: public_learning_call(call)
        for title, call in data.get("learning_calls", {}).items()
        if title in LEARNING_CALL_TITLES and isinstance(call, dict)
    }
    return jsonify({"calls": calls})


@talk_library_bp.get("/tracks")
@login_required
def get_talk_tracks():
    data = load_talk_library()
    normalized_tracks = normalized_talk_tracks(data)
    if not can_manage_accounts():
        normalized_tracks = [track for track in normalized_tracks if track.get("status") != "停用"]
    tracks = [
        public_talk_track(track)
        for track in normalized_tracks
    ]
    return jsonify({"tracks": [track for track in tracks if track is not None]})


@talk_library_bp.get("/my-tracks")
@login_required
def get_my_talk_tracks():
    data = load_talk_library()
    username = str(g.user.get("username") or "").strip()
    tracks = [
        public_personal_talk_track(track)
        for track in normalized_personal_talk_tracks(data, username)
    ]
    return jsonify({"tracks": [track for track in tracks if track is not None]})


@talk_library_bp.post("/my-tracks")
@login_required
def create_my_talk_track():
    payload = request.get_json(silent=True) or {}
    track = normalize_personal_talk_track(payload)
    if track is None:
        return jsonify({"error": "请填写常用话术内容。"}), 400

    now = datetime.now(timezone.utc).isoformat()
    username = str(g.user.get("username") or "").strip()
    track["id"] = uuid.uuid4().hex
    track["created_by"] = username
    track["created_at"] = now
    track["updated_at"] = now

    data = load_talk_library()
    tracks = normalized_personal_talk_tracks(data, username)
    tracks.insert(0, track)
    data.setdefault("personal_tracks", {})[username] = tracks
    save_talk_library(data)
    return jsonify({
        "track": public_personal_talk_track(track),
        "tracks": [public_personal_talk_track(item) for item in normalized_personal_talk_tracks(data, username)],
    }), 201


@talk_library_bp.delete("/my-tracks/<track_id>")
@login_required
def delete_my_talk_track(track_id):
    username = str(g.user.get("username") or "").strip()
    data = load_talk_library()
    tracks = normalized_personal_talk_tracks(data, username)
    if not any(track.get("id") == track_id for track in tracks):
        return jsonify({"error": "常用话术不存在。"}), 404

    data.setdefault("personal_tracks", {})[username] = [
        track for track in tracks if track.get("id") != track_id
    ]
    save_talk_library(data)
    return jsonify({
        "ok": True,
        "tracks": [public_personal_talk_track(item) for item in normalized_personal_talk_tracks(data, username)],
    })


@talk_library_bp.post("/tracks")
@login_required
def create_talk_track():
    payload = request.get_json(silent=True) or {}
    data = load_talk_library()
    track = normalize_talk_track(payload, category_names=set(talk_category_names(data)), require_fields=True)
    if track is None:
        return jsonify({"error": "请填写分类、关键词和话术内容。"}), 400
    if not can_create_talk_track_item(track):
        return jsonify({"error": "只有Joanna账号可以新增该专题话术。"}), 403

    now = datetime.now(timezone.utc).isoformat()
    track["id"] = uuid.uuid4().hex
    track["created_by"] = str(g.user.get("username") or "")
    track["created_at"] = now
    track["updated_at"] = now

    tracks = normalized_talk_tracks(data)
    tracks.insert(0, track)
    data["tracks"] = tracks
    save_talk_library(data)
    return jsonify({
        "track": public_talk_track(track),
        "tracks": [public_talk_track(item) for item in normalized_talk_tracks(data)],
    })


@talk_library_bp.post("/tracks/import")
@login_required
def import_talk_tracks():
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以导入话术。"}), 403

    payload = request.get_json(silent=True) or {}
    raw_tracks = payload.get("tracks") if isinstance(payload.get("tracks"), list) else []
    data = load_talk_library()
    category_names = set(talk_category_names(data))
    replace_category = str(payload.get("category") or "").strip()
    if replace_category not in category_names:
        replace_category = ""
    should_replace_category = bool(payload.get("replace_category"))

    imported, invalid_rows = import_talk_track_rows(data, raw_tracks, replace_category, should_replace_category)
    if invalid_rows:
        return jsonify({
            "error": "分类不存在，请先在工作集里创建或修正分类名" if any("分类不存在" in "；".join(row.get("errors", [])) for row in invalid_rows) else "导入失败，请检查必填字段。",
            "invalid_rows": invalid_rows,
        }), 400

    save_talk_library(data)
    return jsonify({
        "imported_count": len(imported),
        "tracks": [public_talk_track(item) for item in normalized_talk_tracks(data)],
    })


@talk_library_bp.post("/tracks/import-file")
@login_required
def import_talk_tracks_file():
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以导入话术。"}), 403

    uploaded_file = request.files.get("file")
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({"error": "请选择要导入的 Excel 或 CSV 文件。"}), 400
    try:
        raw_tracks = parse_uploaded_talk_file(uploaded_file)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    data = load_talk_library()
    replace_category = str(request.form.get("category") or "").strip()
    should_replace_category = request.form.get("replace_category") in {"1", "true", "True", "on"}
    imported, invalid_rows = import_talk_track_rows(data, raw_tracks, replace_category, should_replace_category)
    if invalid_rows:
        return jsonify({
            "error": "分类不存在，请先在工作集里创建或修正分类名" if any("分类不存在" in "；".join(row.get("errors", [])) for row in invalid_rows) else "导入失败，请检查必填字段。",
            "invalid_rows": invalid_rows,
        }), 400

    save_talk_library(data)
    return jsonify({
        "imported_count": len(imported),
        "tracks": [public_talk_track(item) for item in normalized_talk_tracks(data)],
    })


@talk_library_bp.delete("/tracks/<track_id>")
@login_required
def delete_talk_track(track_id):
    data = load_talk_library()
    tracks = normalized_talk_tracks(data)
    target_track = next((track for track in tracks if track.get("id") == track_id), None)
    if target_track is None:
        return jsonify({"error": "话术不存在。"}), 404
    if not can_delete_talk_track_item(target_track):
        return jsonify({"error": "只能删除自己添加的答疑话术。"}), 403

    next_tracks = [track for track in tracks if track.get("id") != track_id]
    data["tracks"] = next_tracks
    save_talk_library(data)
    return jsonify({"ok": True, "tracks": [public_talk_track(item) for item in normalized_talk_tracks(data)]})


@talk_library_bp.put("/learning-calls")
@login_required
def save_learning_call():
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以维护学情电话话术。"}), 403

    payload = request.get_json(silent=True) or {}
    call = normalize_learning_call(payload)
    if call is None:
        return jsonify({"error": "请选择正确的通话类型。"}), 400
    if not any(call[key] for key in LEARNING_SECTION_KEYS):
        return jsonify({"error": "请至少填写一项学情电话内容。"}), 400

    data = load_talk_library()
    data.setdefault("learning_calls", {})[call["title"]] = call
    save_talk_library(data)
    return jsonify({"call": public_learning_call(call)})


@talk_library_bp.get("/materials")
@login_required
def get_talk_materials():
    data = load_talk_library()
    materials = [
        public_material(material)
        for material in normalized_materials(data)
    ]
    return jsonify({"materials": [material for material in materials if material is not None]})


@talk_library_bp.post("/materials")
@login_required
def upload_talk_material():
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以上传话术素材。"}), 403

    uploaded_file = request.files.get("file")
    if uploaded_file is None or not uploaded_file.filename:
        return jsonify({"error": "请选择要上传的素材图片。"}), 400

    raw_filename = str(uploaded_file.filename or "").strip()
    safe_filename = secure_filename(raw_filename)
    extension = Path(raw_filename).suffix.lower() or Path(safe_filename).suffix.lower()
    if extension not in TALK_MATERIAL_EXTENSIONS:
        return jsonify({"error": "目前仅支持 png、jpg、jpeg、webp、gif 图片素材。"}), 400

    material_id = uuid.uuid4().hex
    filename = f"{material_id}{extension}"
    material_path = talk_material_dir()
    material_path.mkdir(parents=True, exist_ok=True)
    target_path = material_path / filename
    uploaded_file.save(target_path)

    material = {
        "id": material_id,
        "title": str(request.form.get("title") or "").strip()[:80] or Path(raw_filename).stem[:80] or "素材",
        "category": "",
        "keywords": str(request.form.get("keywords") or "").strip()[:160],
        "note": "",
        "filename": filename,
        "original_filename": raw_filename[:160] or f"material{extension}",
        "mime_type": uploaded_file.mimetype or "",
        "size": target_path.stat().st_size if target_path.exists() else 0,
        "uploaded_by": str(g.user.get("username") or ""),
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
    }

    data = load_talk_library()
    data.setdefault("materials", []).append(material)
    save_talk_library(data)
    return jsonify({"material": public_material(material)})


@talk_library_bp.get("/materials/<material_id>/file")
@login_required
def get_talk_material_file(material_id):
    data = load_talk_library()
    material = next(
        (item for item in normalized_materials(data) if item.get("id") == material_id),
        None,
    )
    if material is None:
        return jsonify({"error": "素材不存在。"}), 404
    return send_from_directory(talk_material_dir(), material["filename"])


@talk_library_bp.delete("/materials/<material_id>")
@login_required
def delete_talk_material(material_id):
    if not can_manage_accounts():
        return jsonify({"error": "只有Joanna账号可以删除话术素材。"}), 403

    data = load_talk_library()
    materials = normalized_materials(data)
    material = next((item for item in materials if item.get("id") == material_id), None)
    if material is None:
        return jsonify({"error": "素材不存在。"}), 404

    data["materials"] = [item for item in materials if item.get("id") != material_id]
    save_talk_library(data)

    target_path = talk_material_dir() / material["filename"]
    if target_path.exists():
        target_path.unlink()
    return jsonify({"ok": True})
