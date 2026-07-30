import csv
import io
import re
import uuid
from datetime import date, datetime, timedelta

from flask import Blueprint, g, jsonify, request
from openpyxl import load_workbook

from app.auth import can_manage_accounts, current_teacher_id as auth_current_teacher_id, login_required
from app.classes import (
    LEARNING_BOOKS,
    active_completion_activity,
    can_read_class,
    class_teacher_id,
    completion_period_from_store,
    current_title_week_number,
    current_month_key,
    export_cycle_anchor,
    find_class_by_id,
    get_student_account,
    local_date_key,
    load_activity_store,
    load_store,
    merge_learning_assessments,
    normalize_header,
    normalize_identity,
    now_iso,
    parse_chinese_number,
    parse_upload,
    save_store,
)
from app.database import load_database_settings, save_database_settings
from app.teachers import teacher_label


learning_coaching_bp = Blueprint("learning_coaching", __name__, url_prefix="/api/learning-coaching")

LEARNING_SCORE_CATEGORIES = [
    {"key": "excellent", "label": "培优", "min": 88, "max": 100},
    {"key": "improve", "label": "提升", "min": 72, "max": 88},
    {"key": "good", "label": "良好", "min": 60, "max": 72},
    {"key": "basic", "label": "基础", "min": 0, "max": 60},
]
DEFAULT_COACHING_ROUNDS = [
    {"id": "round-1", "name": "第一轮辅导", "start_week": 15, "deadline_week": 24},
    {"id": "round-2", "name": "第二轮辅导", "start_week": 30, "deadline_week": 39},
]
GUIDE_TEXT_FIELDS = ("words", "sentences", "grammar", "practice", "answer")
GUIDE_FIELD_LABELS = {
    "words": "重点单词",
    "sentences": "重点句型",
    "grammar": "重点语法",
    "practice": "培优题目",
    "answer": "答案",
}


def stage_range_label(stage):
    start_unit = ((int(stage) - 1) * 3) + 1
    return f"Unit{start_unit}-Unit{start_unit + 2}"


def stage_guide_key(book, stage):
    return f"{book}:{int(stage)}"


def guide_key(kind, book, number):
    return f"{kind}:{book}:{int(number)}"


def guide_stage_for_unit(unit):
    return max(1, min(3, ((int(unit) - 1) // 3) + 1))


def guide_label(kind, number):
    return f"Unit{int(number)}" if kind == "unit" else f"阶段{int(number)}"


def normalize_stage_guide(book, stage, item=None):
    source = item if isinstance(item, dict) else {}
    return {
        "book": book,
        "book_label": LEARNING_BOOKS.get(book, "上册"),
        "stage": int(stage),
        "stage_label": f"阶段{int(stage)}",
        "range_label": stage_range_label(stage),
        "focus": str(source.get("focus") or "").strip(),
        "practice": str(source.get("practice") or "").strip(),
        "updated_at": str(source.get("updated_at") or ""),
    }


def normalize_learning_guide(kind, book, number, item=None):
    source = item if isinstance(item, dict) else {}
    safe_kind = "stage" if kind == "stage" else "unit"
    safe_number = int(number)
    payload = {
        "kind": safe_kind,
        "book": book,
        "book_label": LEARNING_BOOKS.get(book, "上册"),
        "number": safe_number,
        "label": guide_label(safe_kind, safe_number),
        "updated_at": str(source.get("updated_at") or ""),
    }
    if safe_kind == "unit":
        payload.update({
            "unit": safe_number,
            "stage": guide_stage_for_unit(safe_number),
            "stage_label": f"阶段{guide_stage_for_unit(safe_number)}",
            "range_label": f"Unit{safe_number}",
        })
    else:
        payload.update({
            "stage": safe_number,
            "stage_label": f"阶段{safe_number}",
            "range_label": stage_range_label(safe_number),
        })
    for field in GUIDE_TEXT_FIELDS:
        payload[field] = str(source.get(field) or "").strip()
    return payload


def load_learning_guides():
    settings = load_database_settings()
    coaching = settings.setdefault("learning_coaching", {})
    guides = coaching.setdefault("guides", {})
    legacy_stage_guides = coaching.setdefault("stage_guides", {})
    normalized = {}
    for book in LEARNING_BOOKS:
        for unit in range(1, 10):
            key = guide_key("unit", book, unit)
            normalized[key] = normalize_learning_guide("unit", book, unit, guides.get(key))
        for stage in range(1, 4):
            key = guide_key("stage", book, stage)
            source = guides.get(key)
            if source is None:
                legacy = legacy_stage_guides.get(stage_guide_key(book, stage))
                if isinstance(legacy, dict):
                    source = {
                        "grammar": legacy.get("focus", ""),
                        "practice": legacy.get("practice", ""),
                        "updated_at": legacy.get("updated_at", ""),
                    }
            normalized[key] = normalize_learning_guide("stage", book, stage, source)
    return settings, guides, normalized


def public_learning_guides():
    _, _, guides = load_learning_guides()
    return guides


def normalize_learning_round(item=None, index=0):
    source = item if isinstance(item, dict) else {}
    try:
        start_week = int(source.get("start_week") or source.get("start") or 0)
    except (TypeError, ValueError):
        start_week = 0
    try:
        deadline_week = int(source.get("deadline_week") or source.get("end_week") or source.get("deadline") or 0)
    except (TypeError, ValueError):
        deadline_week = 0
    start_week = max(1, min(99, start_week))
    deadline_week = max(start_week, min(99, deadline_week or start_week))
    name = str(source.get("name") or f"第{index + 1}轮辅导").strip()[:20]
    round_id = str(source.get("id") or f"round-{index + 1}").strip()[:40]
    return {
        "id": round_id or f"round-{index + 1}",
        "name": name or f"第{index + 1}轮辅导",
        "start_week": start_week,
        "deadline_week": deadline_week,
    }


def default_learning_rounds():
    return [
        normalize_learning_round(item, index)
        for index, item in enumerate(DEFAULT_COACHING_ROUNDS)
    ]


def normalize_learning_rounds(items):
    source = items if isinstance(items, list) else []
    rounds = []
    seen = set()
    for index, item in enumerate(source):
        normalized = normalize_learning_round(item, index)
        if not normalized["start_week"]:
            continue
        if normalized["id"] in seen:
            normalized["id"] = f"{normalized['id']}-{index + 1}"
        seen.add(normalized["id"])
        rounds.append(normalized)
    rounds = sorted(rounds, key=lambda row: (row["start_week"], row["deadline_week"], row["name"]))
    return rounds or default_learning_rounds()


def load_learning_rounds():
    settings = load_database_settings()
    coaching = settings.setdefault("learning_coaching", {})
    raw_rounds = coaching.get("rounds")
    rounds = normalize_learning_rounds(raw_rounds)
    if raw_rounds != rounds:
        coaching["rounds"] = rounds
        save_database_settings(settings)
    return settings, rounds


def public_learning_rounds():
    _, rounds = load_learning_rounds()
    return rounds


def learning_today_key():
    return datetime.now().date().isoformat()


def parse_appointment_date(value):
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError:
        return ""


def learning_appointment_key(class_id, student_id, book, kind, number, round_id=""):
    safe_kind = "stage" if kind == "stage" else "unit"
    safe_round_id = str(round_id or "").strip()
    if safe_round_id:
        return f"{class_id}:{student_id}:{book}:{safe_round_id}:{safe_kind}:{int(number or 1)}"
    return f"{class_id}:{student_id}:{book}:{safe_kind}:{int(number or 1)}"


def normalize_learning_appointment(key, item=None):
    source = item if isinstance(item, dict) else {}
    parts = str(key or "").split(":")
    class_id = str(source.get("class_id") or (parts[0] if len(parts) >= 5 else "")).strip()
    student_id = str(source.get("student_id") or (parts[1] if len(parts) >= 5 else "")).strip()
    book = str(source.get("book") or (parts[2] if len(parts) >= 5 else "upper")).strip()
    if book not in LEARNING_BOOKS:
        book = "upper"
    round_id = str(source.get("round_id") or "").strip()
    if len(parts) >= 6:
        round_id = round_id or parts[3]
        kind_part = parts[4]
        number_part = parts[5]
    else:
        kind_part = parts[3] if len(parts) >= 5 else "unit"
        number_part = parts[4] if len(parts) >= 5 else 1
    kind = str(source.get("kind") or kind_part).strip()
    kind = "stage" if kind == "stage" else "unit"
    try:
        number = int(source.get("number") or number_part)
    except (TypeError, ValueError):
        number = 1
    number = max(1, min(3 if kind == "stage" else 9, number))
    safe_key = learning_appointment_key(class_id, student_id, book, kind, number, round_id)
    appointment_date = parse_appointment_date(source.get("appointment_date"))
    return {
        "key": safe_key,
        "class_id": class_id,
        "student_id": student_id,
        "book": book,
        "round_id": round_id,
        "kind": kind,
        "number": number,
        "appointment_date": appointment_date,
        "completed": bool(source.get("completed")),
        "completed_at": str(source.get("completed_at") or ""),
        "completed_date": parse_appointment_date(source.get("completed_date")),
        "created_at": str(source.get("created_at") or ""),
        "updated_at": str(source.get("updated_at") or ""),
    }


def load_learning_appointments():
    settings = load_database_settings()
    coaching = settings.setdefault("learning_coaching", {})
    appointments = coaching.setdefault("appointments", {})
    if not isinstance(appointments, dict):
        appointments = {}
        coaching["appointments"] = appointments
    normalized = {}
    for raw_key, raw_item in appointments.items():
        item = normalize_learning_appointment(raw_key, raw_item)
        if item["class_id"] and item["student_id"]:
            normalized[item["key"]] = item
    coaching["appointments"] = normalized
    return settings, coaching["appointments"], normalized


def find_student_by_id(target_class, student_id):
    for student in target_class.get("students", []):
        if str(student.get("id") or "") == str(student_id or ""):
            return student
    return None


def appointment_target_label(item):
    if item.get("kind") == "stage":
        return f"{stage_range_label(item.get('number') or 1)} 阶段测评"
    return f"Unit{int(item.get('number') or 1)} 单元检测"


def appointment_date_label(date_key):
    parsed = parse_appointment_date(date_key)
    if not parsed:
        return ""
    return learning_date_label(date.fromisoformat(parsed))


def public_learning_appointments(appointments, class_items):
    today = learning_today_key()
    class_map = {str(item.get("id") or ""): item for item in class_items}
    public_items = {}
    due_items = []
    for raw_item in appointments.values():
        item = normalize_learning_appointment(raw_item.get("key"), raw_item)
        if not item.get("appointment_date") and not item.get("completed"):
            continue
        target_class = class_map.get(item["class_id"])
        if not target_class:
            continue
        student = find_student_by_id(target_class, item["student_id"])
        if not student or student.get("learning_coaching_hidden"):
            continue
        payload = {
            **item,
            "class_name": target_class.get("name", ""),
            "teacher_id": class_teacher_id(target_class),
            "teacher_name": teacher_label(class_teacher_id(target_class)),
            "student_name": str(student.get("name") or "").strip(),
            "student_account": get_student_account(student),
            "book_label": LEARNING_BOOKS.get(item["book"], ""),
            "round_id": item.get("round_id", ""),
            "target_label": appointment_target_label(item),
            "appointment_label": appointment_date_label(item["appointment_date"]),
            "completed_label": appointment_date_label(item["completed_date"]),
            "is_due": item["appointment_date"] <= today,
            "is_today": item["appointment_date"] == today,
        }
        public_items[payload["key"]] = payload
        if not payload["completed"] and payload["is_due"]:
            due_items.append(payload)
    due_items = sorted(
        due_items,
        key=lambda row: (
            row.get("appointment_date") or "9999-12-31",
            row.get("teacher_name", ""),
            row.get("class_name", ""),
            row.get("student_name", ""),
            row.get("target_label", ""),
        ),
    )
    return public_items, due_items


def learning_appointment_payload(store=None):
    source = store or load_store()
    class_items = [item for item in source.get("classes", []) if can_read_class(item)]
    _, _, appointments = load_learning_appointments()
    public_items, due_items = public_learning_appointments(appointments, class_items)
    return {
        "appointments": public_items,
        "today_appointments": due_items,
    }


def appointment_request_context(payload):
    store = load_store()
    class_id = str(payload.get("class_id") or "").strip()
    target_class = find_class_by_id(store, class_id)
    if target_class is None or not can_read_class(target_class):
        return None, None, {"error": "班级不存在。", "status": 404}
    if not can_write_learning_class(target_class):
        return None, None, {"error": "只能维护自己班级的辅导预约。", "status": 403}
    student_id = str(payload.get("student_id") or "").strip()
    student = find_student_by_id(target_class, student_id)
    if not student or student.get("learning_coaching_hidden"):
        return None, None, {"error": "学员不存在。", "status": 404}
    book = str(payload.get("book") or "upper").strip()
    if book not in LEARNING_BOOKS:
        return None, None, {"error": "册别不存在。", "status": 400}
    kind = "stage" if str(payload.get("kind") or "") == "stage" else "unit"
    try:
        number = int(payload.get("number") or 1)
    except (TypeError, ValueError):
        number = 1
    if kind == "unit" and (number < 1 or number > 9):
        return None, None, {"error": "单元不存在。", "status": 400}
    if kind == "stage" and (number < 1 or number > 3):
        return None, None, {"error": "阶段不存在。", "status": 400}
    round_id = str(payload.get("round_id") or "").strip()
    return store, {
        "class_id": class_id,
        "student_id": student_id,
        "book": book,
        "round_id": round_id,
        "kind": kind,
        "number": number,
    }, None


def guide_header_key(value):
    normalized = normalize_header(value)
    aliases = {
        "book": {"册别", "上下册", "教材", "pu", "pu册别"},
        "kind": {"类型", "内容类型", "资料类型", "类别"},
        "number": {"编号", "单元", "阶段", "单元阶段", "范围", "unit", "unit阶段"},
        "words": {"重点单词", "单词", "核心单词", "词汇"},
        "sentences": {"重点句型", "句型", "核心句型"},
        "grammar": {"重点语法", "语法", "核心语法"},
        "practice": {"培优题目", "题目", "培优练习", "练习题", "拓展题"},
        "answer": {"答案", "参考答案", "解析", "题目答案"},
    }
    for key, names in aliases.items():
        if normalized in {normalize_header(name) for name in names}:
            return key
    return ""


def guide_rows_from_upload(file_storage):
    filename = str(file_storage.filename or "").lower()
    raw_rows = []
    if filename.endswith(".csv"):
        content = file_storage.read().decode("utf-8-sig")
        raw_rows = [row for row in csv.reader(io.StringIO(content))]
    elif filename.endswith(".xlsx"):
        workbook = load_workbook(file_storage, data_only=True, read_only=True)
        try:
            sheet = workbook.active
            raw_rows = [
                ["" if cell is None else cell for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    else:
        raise ValueError("仅支持 .xlsx 或 .csv 资料表。")
    return raw_rows


def guide_cell(row, indexes, key):
    index = indexes.get(key)
    if index is None or index >= len(row):
        return ""
    value = row[index]
    return "" if value is None else str(value).strip()


def parse_guide_book(*values):
    text = " ".join(str(value or "") for value in values).lower()
    compact = re.sub(r"\s+", "", text)
    if "pu2" in compact or "下册" in text or "下半册" in text:
        return "lower"
    if "pu1" in compact or "上册" in text or "上半册" in text:
        return "upper"
    return ""


def parse_guide_number(value, kind):
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"unit\s*([0-9]{1,2})\s*[-~至到]\s*([0-9]{1,2})", text, flags=re.IGNORECASE)
    if match:
        start = int(match.group(1))
        return guide_stage_for_unit(start) if kind == "stage" else start
    match = re.search(r"([0-9]{1,2})\s*[-~至到]\s*([0-9]{1,2})", text)
    if match:
        start = int(match.group(1))
        return guide_stage_for_unit(start) if kind == "stage" else start
    match = re.search(r"unit\s*([0-9]{1,2})", text, flags=re.IGNORECASE)
    if match:
        number = int(match.group(1))
        return guide_stage_for_unit(number) if kind == "stage" else number
    match = re.search(r"[第]?\s*([0-9一二两三四五六七八九十]{1,3})\s*(?:单元|阶段)?", text)
    if match:
        number = parse_chinese_number(match.group(1))
        if number:
            return number
    return None


def parse_guide_kind(kind_value, number_value):
    text = f"{kind_value or ''} {number_value or ''}"
    if "阶段" in text or re.search(r"[0-9]\s*[-~至到]\s*[0-9]", text) or re.search(r"unit\s*[0-9]{1,2}\s*[-~至到]", text, flags=re.IGNORECASE):
        return "stage"
    return "unit"


def parse_learning_guide_upload(file_storage):
    rows = guide_rows_from_upload(file_storage)
    header_index = None
    indexes = {}
    for row_index, row in enumerate(rows[:12]):
        candidate = {}
        for index, value in enumerate(row):
            key = guide_header_key(value)
            if key and key not in candidate:
                candidate[key] = index
        if {"number", "words", "sentences", "grammar", "practice", "answer"} & set(candidate):
            header_index = row_index
            indexes = candidate
            break
    if header_index is None:
        raise ValueError("未识别到资料表表头，请至少包含：册别、类型、编号、重点单词、重点句型、重点语法、培优题目、答案。")

    parsed = []
    for row in rows[header_index + 1:]:
        if not any(str(value or "").strip() for value in row):
            continue
        raw_book = guide_cell(row, indexes, "book")
        raw_kind = guide_cell(row, indexes, "kind")
        raw_number = guide_cell(row, indexes, "number")
        row_text = " ".join(str(value or "") for value in row)
        book = parse_guide_book(raw_book, raw_number, row_text)
        if not book:
            continue
        kind = parse_guide_kind(raw_kind, raw_number)
        number = parse_guide_number(raw_number or row_text, kind)
        if kind == "unit" and (number is None or number < 1 or number > 9):
            continue
        if kind == "stage" and (number is None or number < 1 or number > 3):
            continue
        data = {
            field: guide_cell(row, indexes, field)
            for field in GUIDE_TEXT_FIELDS
        }
        if not any(data.values()):
            continue
        parsed.append(normalize_learning_guide(kind, book, number, {
            **data,
            "updated_at": now_iso(),
        }))
    if not parsed:
        raise ValueError("资料表里没有识别到可导入的内容，请检查册别、类型和编号。")
    return parsed


def load_stage_guides():
    settings = load_database_settings()
    coaching = settings.setdefault("learning_coaching", {})
    guides = coaching.setdefault("stage_guides", {})
    normalized = {}
    for book in LEARNING_BOOKS:
        for stage in range(1, 4):
            key = stage_guide_key(book, stage)
            normalized[key] = normalize_stage_guide(book, stage, guides.get(key))
    return settings, guides, normalized


def public_stage_guides():
    _, _, guides = load_stage_guides()
    return guides


def can_write_learning_class(item):
    return item.get("owner") == g.user["username"] or can_manage_accounts()


def learning_roster_students(target_class):
    return [
        student
        for student in target_class.get("students", [])
        if not student.get("learning_coaching_hidden")
    ]


def learning_roster_accounts(target_class):
    roster = target_class.get("learning_coaching_roster")
    source = roster.get("accounts") if isinstance(roster, dict) else None
    if not isinstance(source, list):
        return []
    accounts = []
    seen = set()
    for account in source:
        normalized = normalize_identity(account)
        if normalized and normalized not in seen:
            seen.add(normalized)
            accounts.append(normalized)
    return accounts


def imported_learning_roster_accounts(imported_students):
    accounts = []
    seen = set()
    for imported in imported_students:
        normalized = normalize_identity(imported.get("account"))
        if normalized and normalized not in seen:
            seen.add(normalized)
            accounts.append(normalized)
    return accounts


def learning_roster_display_names(imported_students, account_keys):
    wanted = set(account_keys)
    labels = []
    for imported in imported_students:
        account = normalize_identity(imported.get("account"))
        if account not in wanted:
            continue
        name = str(imported.get("name") or "").strip()
        raw_account = str(imported.get("account") or "").strip()
        labels.append(f"{name or '未命名学员'}（{raw_account or account}）")
        if len(labels) >= 5:
            break
    return labels


def prepare_learning_roster(target_class, imported_students, updated_at):
    incoming_accounts = imported_learning_roster_accounts(imported_students)
    if not incoming_accounts:
        return {
            "error": "没有识别到学习账号，无法建立辅导名单。",
            "status": 400,
        }

    existing_accounts = learning_roster_accounts(target_class)
    incoming_set = set(incoming_accounts)
    existing_set = set(existing_accounts)
    is_initial = not existing_accounts

    if existing_accounts:
        extra_accounts = incoming_set - existing_set
        if extra_accounts:
            labels = learning_roster_display_names(imported_students, extra_accounts)
            preview = "、".join(labels) or "未知账号"
            suffix = "等" if len(extra_accounts) > len(labels) else ""
            return {
                "error": f"本次表格发现 {len(extra_accounts)} 个不在首次辅导名单中的账号：{preview}{suffix}。请确认班级后再上传。",
                "status": 400,
            }

    active_accounts = incoming_accounts if is_initial else [
        account for account in existing_accounts if account in incoming_set
    ]
    active_set = set(active_accounts)
    removed_accounts = existing_set - active_set if existing_accounts else set()
    hidden_student_ids = set()

    for student in target_class.get("students", []):
        account = normalize_identity(get_student_account(student))
        if not account:
            continue
        if account in active_set:
            continue
        student["learning_coaching_hidden"] = True
        student["updated_at"] = updated_at
        student_id = str(student.get("id") or "").strip()
        if student_id:
            hidden_student_ids.add(student_id)

    target_class["learning_coaching_roster"] = {
        "accounts": active_accounts,
        "initialized_at": (
            target_class.get("learning_coaching_roster", {}).get("initialized_at")
            if isinstance(target_class.get("learning_coaching_roster"), dict)
            else ""
        ) or updated_at,
        "updated_at": updated_at,
    }
    return {
        "accounts": active_set,
        "is_initial": is_initial,
        "removed_accounts": removed_accounts,
        "hidden_student_ids": hidden_student_ids,
    }


def score_category(score):
    try:
        value = float(score)
    except (TypeError, ValueError):
        return {"key": "missing", "label": "暂无分数"}
    if value >= 88:
        return {"key": "excellent", "label": "培优"}
    if value >= 72:
        return {"key": "improve", "label": "提升"}
    if value >= 60:
        return {"key": "good", "label": "良好"}
    return {"key": "basic", "label": "基础"}


def normalized_score(value):
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def assessment_score_payload(assessment):
    score = normalized_score(assessment.get("score")) if assessment else None
    if score is None:
        return {
            "score": None,
            "category": "",
            "category_key": "missing",
            "source_label": "",
            "updated_at": "",
        }
    category = score_category(score)
    return {
        "score": score,
        "category": category["label"],
        "category_key": category["key"],
        "source_label": assessment.get("label", ""),
        "updated_at": assessment.get("updated_at", ""),
    }


def assessment_matches(item, book, assessment_type, unit=None, stage=None):
    if not isinstance(item, dict):
        return False
    if str(item.get("book") or "upper") != book:
        return False
    if str(item.get("type") or "unit") != assessment_type:
        return False
    if assessment_type == "stage":
        return int(item.get("stage") or 0) == int(stage or 0)
    return int(item.get("unit") or 0) == int(unit or 0)


def latest_assessment(student, book, assessment_type, unit=None, stage=None):
    matches = [
        item
        for item in student.get("learning_assessments", [])
        if assessment_matches(item, book, assessment_type, unit, stage)
    ]
    if not matches:
        return None
    return max(matches, key=lambda item: str(item.get("updated_at") or ""))


def assessment_students(target_class, book, assessment_type, unit=None, stage=None):
    rows = []
    for student in learning_roster_students(target_class):
        assessment = latest_assessment(student, book, assessment_type, unit, stage)
        if not assessment:
            continue
        score = normalized_score(assessment.get("score"))
        if score is None:
            continue
        category = score_category(score)
        rows.append({
            "student_id": student.get("id", ""),
            "name": str(student.get("name") or "").strip(),
            "account": get_student_account(student),
            "score": score,
            "category": category["label"],
            "category_key": category["key"],
            "source_label": assessment.get("label", ""),
            "updated_at": assessment.get("updated_at", ""),
        })
    return sorted(rows, key=lambda item: (-item["score"], item["name"], item["account"]))


def assessment_summary(target_class, book, assessment_type, unit=None, stage=None):
    rows = assessment_students(target_class, book, assessment_type, unit, stage)
    counts = {item["key"]: 0 for item in LEARNING_SCORE_CATEGORIES}
    for row in rows:
        counts[row["category_key"]] = counts.get(row["category_key"], 0) + 1
    scores = [row["score"] for row in rows]
    average_score = round(sum(scores) / len(scores), 2) if scores else None
    return {
        "count": len(rows),
        "student_total": len(learning_roster_students(target_class)),
        "average_score": average_score,
        "counts": counts,
        "students": rows,
    }


def unit_payload(target_class, book, unit):
    summary = assessment_summary(target_class, book, "unit", unit=unit)
    return {
        "type": "unit",
        "book": book,
        "unit": unit,
        "label": f"Unit{unit}",
        **summary,
    }


def stage_payload(target_class, book, stage):
    summary = assessment_summary(target_class, book, "stage", stage=stage)
    return {
        "type": "stage",
        "book": book,
        "stage": stage,
        "label": f"阶段{stage}",
        "range_label": stage_range_label(stage),
        **summary,
    }


def unit_score_matrix_rows(target_class, book):
    rows = []
    for student in learning_roster_students(target_class):
        unit_scores = []
        scored_count = 0
        score_values = []
        stage_assessments = {
            stage: latest_assessment(student, book, "stage", stage=stage)
            for stage in range(1, 4)
        }
        for unit in range(1, 10):
            stage = max(1, min(3, ((unit - 1) // 3) + 1))
            assessment = latest_assessment(student, book, "unit", unit=unit)
            score = normalized_score(assessment.get("score")) if assessment else None
            if score is None:
                unit_scores.append({
                    "unit": unit,
                    "stage": stage,
                    "stage_label": f"阶段{stage}",
                    "stage_range_label": stage_range_label(stage),
                    "score": None,
                    "category": "",
                    "category_key": "missing",
                    "source_label": "",
                    "updated_at": "",
                    "stage_score": assessment_score_payload(stage_assessments.get(stage)),
                })
                continue
            category = score_category(score)
            scored_count += 1
            score_values.append(score)
            unit_scores.append({
                "unit": unit,
                "stage": stage,
                "stage_label": f"阶段{stage}",
                "stage_range_label": stage_range_label(stage),
                "score": score,
                "category": category["label"],
                "category_key": category["key"],
                "source_label": assessment.get("label", ""),
                "updated_at": assessment.get("updated_at", ""),
                "stage_score": assessment_score_payload(stage_assessments.get(stage)),
            })
        average_score = round(sum(score_values) / len(score_values), 2) if score_values else None
        rows.append({
            "student_id": student.get("id", ""),
            "name": str(student.get("name") or "").strip(),
            "account": get_student_account(student),
            "scored_count": scored_count,
            "average_score": average_score,
            "unit_scores": unit_scores,
        })
    return sorted(
        rows,
        key=lambda item: (
            -int(item.get("scored_count") or 0),
            -(item.get("average_score") if item.get("average_score") is not None else -1),
            item.get("name", ""),
            item.get("account", ""),
        ),
    )


def class_learning_payload(item, rounds=None):
    teacher_id = class_teacher_id(item)
    week_number = current_title_week_number(item)
    cycle = learning_cycle_info(week_number, rounds)
    books = {}
    for book_key, book_label in LEARNING_BOOKS.items():
        units = [unit_payload(item, book_key, unit) for unit in range(1, 10)]
        stages = [stage_payload(item, book_key, stage) for stage in range(1, 4)]
        books[book_key] = {
            "key": book_key,
            "label": book_label,
            "units": units,
            "stages": stages,
            "student_rows": unit_score_matrix_rows(item, book_key),
            "uploaded_count": sum(unit["count"] for unit in units) + sum(stage["count"] for stage in stages),
        }
    assessment_count = sum(book["uploaded_count"] for book in books.values())
    return {
        "id": item.get("id", ""),
        "name": item.get("name", ""),
        "note": str(item.get("note") or "").strip(),
        "teacher_id": teacher_id,
        "teacher_name": teacher_label(teacher_id),
        "owner": item.get("owner", ""),
        "student_count": len(learning_roster_students(item)),
        "can_write": can_write_learning_class(item),
        "title_week_number": week_number,
        "title_week_label": f"W{week_number}" if week_number else "",
        "in_coaching_cycle": cycle["status"] == "active",
        "coaching_cycle": cycle,
        "assessment_count": assessment_count,
        "books": books,
    }


def find_learning_student(target_class, imported):
    account_key = normalize_identity(imported.get("account"))
    for student in target_class.get("students", []):
        if account_key and normalize_identity(get_student_account(student)) == account_key:
            return student
    return None


def merge_learning_upload(target_class, imported_students, month_key, updated_at, roster_accounts=None):
    students = target_class.setdefault("students", [])
    created = 0
    updated = 0
    score_updated = 0
    score_count = 0
    skipped = 0
    allowed_accounts = set(roster_accounts or [])
    for imported in imported_students:
        account = str(imported.get("account") or "").strip()
        account_key = normalize_identity(account)
        assessments = imported.get("assessments", [])
        if not account_key:
            skipped += 1
            continue
        if allowed_accounts and account_key not in allowed_accounts:
            skipped += 1
            continue

        student = find_learning_student(target_class, imported)
        if student is None:
            student = {
                "id": imported.get("id") or uuid.uuid4().hex,
                "name": str(imported.get("name") or "").strip(),
                "account": account,
                "months": {},
                "created_at": updated_at,
            }
            students.append(student)
            created += 1
        else:
            updated += 1

        if not student.get("name_locked") and not str(student.get("name") or "").strip():
            student["name"] = str(imported.get("name") or student.get("name") or "").strip()
        student["account"] = account or str(student.get("account") or "").strip()
        student["updated_at"] = updated_at
        if assessments:
            score_count += len(assessments)
            score_updated += merge_learning_assessments(student, assessments, month_key, updated_at)

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "score_count": score_count,
        "assessment_updated": score_updated,
    }


def remove_learning_appointments_for_students(class_id, student_ids):
    safe_ids = {str(student_id or "").strip() for student_id in student_ids if str(student_id or "").strip()}
    if not safe_ids:
        return 0
    settings, appointments, _ = load_learning_appointments()
    retained = {
        key: item
        for key, item in appointments.items()
        if not (
            str(item.get("class_id") or "").strip() == str(class_id or "").strip()
            and str(item.get("student_id") or "").strip() in safe_ids
        )
    }
    removed = len(appointments) - len(retained)
    if removed:
        settings.setdefault("learning_coaching", {})["appointments"] = retained
        save_database_settings(settings)
    return removed


def learning_date_label(value):
    if value is None:
        return ""
    return f"{value.month}月{value.day}日"


def learning_cycle_date_for_week(current_week, target_week):
    if not current_week:
        return None
    week_delta = int(target_week) - int(current_week)
    return export_cycle_anchor() + timedelta(days=week_delta * 7)


def learning_cycle_info(week_number):
    if not week_number:
        return {
            "status": "unknown",
            "status_label": "未填写W数",
            "entry_date": "",
            "entry_label": "",
            "deadline_date": "",
            "deadline_label": "",
            "display_text": "先在我的班级里填写周数",
        }

    entry_date = learning_cycle_date_for_week(week_number, COACHING_START_WEEK)
    deadline_date = learning_cycle_date_for_week(week_number, COACHING_DEADLINE_WEEK)
    entry_label = learning_date_label(entry_date)
    deadline_label = learning_date_label(deadline_date)

    if week_number < COACHING_START_WEEK:
        status = "upcoming"
        status_label = "未进入第一次辅导周期"
        display_text = f"{entry_label}进入第一次辅导周期" if entry_label else "即将进入第一次辅导周期"
    elif week_number <= COACHING_DEADLINE_WEEK:
        status = "active"
        status_label = "第一次辅导周期中"
        display_text = f"第一次辅导周期中 · DDL {deadline_label}" if deadline_label else "第一次辅导周期中"
    else:
        status = "ended"
        status_label = "已过第一次辅导周期"
        display_text = f"已过第一次辅导周期 · DDL {deadline_label}" if deadline_label else "已过第一次辅导周期"

    return {
        "status": status,
        "status_label": status_label,
        "entry_date": local_date_key(entry_date) if entry_date else "",
        "entry_label": entry_label,
        "deadline_date": local_date_key(deadline_date) if deadline_date else "",
        "deadline_label": deadline_label,
        "display_text": display_text,
    }


def learning_date_label(value):
    if value is None:
        return ""
    return f"{value.month}月{value.day}日"


def learning_cycle_info(week_number, rounds=None):
    safe_rounds = normalize_learning_rounds(rounds)
    if not week_number:
        return {
            "status": "unknown",
            "status_label": "未填写W数",
            "round_id": "",
            "round_name": "",
            "start_week": "",
            "deadline_week": "",
            "entry_date": "",
            "entry_label": "",
            "deadline_date": "",
            "deadline_label": "",
            "display_text": "先在我的班级里填写周数",
        }

    current_week = int(week_number)
    active_round = next(
        (item for item in safe_rounds if item["start_week"] <= current_week <= item["deadline_week"]),
        None,
    )
    upcoming_round = next(
        (item for item in safe_rounds if current_week < item["start_week"]),
        None,
    )
    target_round = active_round or upcoming_round or safe_rounds[-1]

    entry_date = learning_cycle_date_for_week(current_week, target_round["start_week"])
    deadline_date = learning_cycle_date_for_week(current_week, target_round["deadline_week"])
    entry_label = learning_date_label(entry_date)
    deadline_label = learning_date_label(deadline_date)

    if active_round:
        status = "active"
        status_label = f"{target_round['name']}周期中"
        display_text = f"{target_round['name']}周期中 · DDL {deadline_label}" if deadline_label else f"{target_round['name']}周期中"
    elif upcoming_round:
        status = "upcoming"
        status_label = f"未进入{target_round['name']}周期"
        display_text = f"{entry_label}进入{target_round['name']}周期" if entry_label else f"即将进入{target_round['name']}周期"
    else:
        status = "ended"
        status_label = f"已过{target_round['name']}周期"
        display_text = f"已过{target_round['name']}周期 · DDL {deadline_label}" if deadline_label else f"已过{target_round['name']}周期"

    return {
        "status": status,
        "status_label": status_label,
        "round_id": target_round["id"],
        "round_name": target_round["name"],
        "start_week": target_round["start_week"],
        "deadline_week": target_round["deadline_week"],
        "entry_date": local_date_key(entry_date) if entry_date else "",
        "entry_label": entry_label,
        "deadline_date": local_date_key(deadline_date) if deadline_date else "",
        "deadline_label": deadline_label,
        "display_text": display_text,
    }


def learning_class_sort_key(item):
    cycle = item.get("coaching_cycle") or {}
    status_order = {
        "active": 0,
        "upcoming": 1,
        "unknown": 2,
        "ended": 3,
    }
    status = cycle.get("status") or ""
    primary_date = cycle.get("deadline_date") if status == "active" else cycle.get("entry_date")
    return (
        status_order.get(status, 4),
        primary_date or "9999-12-31",
        item.get("teacher_name", ""),
        item.get("title_week_number") or 999,
        item.get("name", ""),
    )


@learning_coaching_bp.get("")
@login_required
def learning_coaching_summary():
    store = load_store()
    _, coaching_rounds = load_learning_rounds()
    class_items = [item for item in store.get("classes", []) if can_read_class(item)]
    classes = [
        class_learning_payload(item, coaching_rounds)
        for item in class_items
    ]
    classes = sorted(
        classes,
        key=learning_class_sort_key,
    )
    appointment_payload = learning_appointment_payload(store)
    return jsonify({
        "classes": classes,
        "books": [{"key": key, "label": label} for key, label in LEARNING_BOOKS.items()],
        "categories": LEARNING_SCORE_CATEGORIES,
        "guides": public_learning_guides(),
        "stage_guides": public_stage_guides(),
        "appointments": appointment_payload["appointments"],
        "today_appointments": appointment_payload["today_appointments"],
        "coaching_rounds": coaching_rounds,
        "can_manage": can_manage_accounts(),
        "current_teacher_id": auth_current_teacher_id(),
        "coaching_start_week": coaching_rounds[0]["start_week"] if coaching_rounds else 15,
        "coaching_deadline_week": coaching_rounds[0]["deadline_week"] if coaching_rounds else 24,
    })


@learning_coaching_bp.put("/rounds")
@login_required
def update_learning_rounds():
    if not can_manage_accounts():
        return jsonify({"error": "只有管理员可以设置辅导周期。"}), 403
    payload = request.get_json(silent=True) or {}
    rounds = normalize_learning_rounds(payload.get("rounds"))
    if not rounds:
        return jsonify({"error": "请至少保留一轮辅导周期。"}), 400
    settings = load_database_settings()
    coaching = settings.setdefault("learning_coaching", {})
    coaching["rounds"] = rounds
    save_database_settings(settings)
    return jsonify({"coaching_rounds": rounds})


@learning_coaching_bp.patch("/<class_id>/students/<student_id>")
@login_required
def update_learning_student(class_id, student_id):
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    if not name:
        return jsonify({"error": "请输入学员姓名。"}), 400

    store = load_store()
    target_class = find_class_by_id(store, class_id)
    if target_class is None or not can_read_class(target_class):
        return jsonify({"error": "班级不存在。"}), 404
    if not can_write_learning_class(target_class):
        return jsonify({"error": "只能维护自己班级的学员姓名。"}), 403

    student = find_student_by_id(target_class, student_id)
    if student is None:
        return jsonify({"error": "学员不存在。"}), 404

    updated_at = now_iso()
    student["name"] = name
    student["name_locked"] = True
    student["updated_at"] = updated_at
    target_class["updated_at"] = updated_at
    save_store(store)
    return jsonify({
        "student": {
            "student_id": student.get("id", ""),
            "name": str(student.get("name") or "").strip(),
            "account": get_student_account(student),
        },
        "class": class_learning_payload(target_class, public_learning_rounds()),
    })


@learning_coaching_bp.put("/appointments")
@login_required
def update_learning_appointment():
    payload = request.get_json(silent=True) or {}
    store, context, error = appointment_request_context(payload)
    if error:
        return jsonify({"error": error["error"]}), error["status"]

    appointment_date = parse_appointment_date(payload.get("appointment_date"))
    settings, appointments, _ = load_learning_appointments()
    key = learning_appointment_key(
        context["class_id"],
        context["student_id"],
        context["book"],
        context["kind"],
        context["number"],
        context.get("round_id", ""),
    )
    if not appointment_date:
        appointments.pop(key, None)
        save_database_settings(settings)
        return jsonify(learning_appointment_payload(store))

    previous = appointments.get(key) if isinstance(appointments.get(key), dict) else {}
    appointments[key] = normalize_learning_appointment(key, {
        **context,
        "appointment_date": appointment_date,
        "completed": False,
        "completed_at": "",
        "completed_date": "",
        "created_at": previous.get("created_at") or now_iso(),
        "updated_at": now_iso(),
    })
    save_database_settings(settings)
    return jsonify(learning_appointment_payload(store))


@learning_coaching_bp.post("/appointments/complete")
@login_required
def complete_learning_appointment():
    payload = request.get_json(silent=True) or {}
    key = str(payload.get("key") or "").strip()
    if key:
        _, _, loaded = load_learning_appointments()
        appointment = loaded.get(key)
        if not appointment:
            return jsonify({"error": "预约记录不存在。"}), 404
        payload = {
            **appointment,
            **payload,
            "class_id": appointment["class_id"],
            "student_id": appointment["student_id"],
            "book": appointment["book"],
            "round_id": appointment.get("round_id", ""),
            "kind": appointment["kind"],
            "number": appointment["number"],
        }

    store, context, error = appointment_request_context(payload)
    if error:
        return jsonify({"error": error["error"]}), error["status"]

    settings, appointments, _ = load_learning_appointments()
    key = learning_appointment_key(
        context["class_id"],
        context["student_id"],
        context["book"],
        context["kind"],
        context["number"],
        context.get("round_id", ""),
    )
    previous = appointments.get(key) if isinstance(appointments.get(key), dict) else {}
    appointments[key] = normalize_learning_appointment(key, {
        **context,
        "appointment_date": parse_appointment_date(previous.get("appointment_date")) or learning_today_key(),
        "completed": True,
        "completed_at": now_iso(),
        "completed_date": learning_today_key(),
        "created_at": previous.get("created_at") or now_iso(),
        "updated_at": now_iso(),
    })
    save_database_settings(settings)
    return jsonify(learning_appointment_payload(store))


@learning_coaching_bp.post("/guides/upload")
@login_required
def upload_learning_guides():
    if not can_manage_accounts():
        return jsonify({"error": "只有管理员可以上传辅导资料。"}), 403
    file_storage = request.files.get("file")
    if not file_storage:
        return jsonify({"error": "请上传 Excel 或 CSV 资料表。"}), 400
    try:
        imported_guides = parse_learning_guide_upload(file_storage)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    settings, guides, _ = load_learning_guides()
    for guide in imported_guides:
        guides[guide_key(guide["kind"], guide["book"], guide["number"])] = guide
    save_database_settings(settings)
    return jsonify({
        "count": len(imported_guides),
        "guides": public_learning_guides(),
    })


@learning_coaching_bp.put("/guides/<kind>/<book>/<int:number>")
@login_required
def update_learning_guide(kind, book, number):
    if not can_manage_accounts():
        return jsonify({"error": "只有管理员可以维护辅导资料。"}), 403
    if kind not in {"unit", "stage"} or book not in LEARNING_BOOKS:
        return jsonify({"error": "辅导资料不存在。"}), 404
    if kind == "unit" and (number < 1 or number > 9):
        return jsonify({"error": "单元资料不存在。"}), 404
    if kind == "stage" and (number < 1 or number > 3):
        return jsonify({"error": "阶段资料不存在。"}), 404

    payload = request.get_json(silent=True) or {}
    guide = normalize_learning_guide(
        kind,
        book,
        number,
        {
            **{field: str(payload.get(field) or "").strip()[:5000] for field in GUIDE_TEXT_FIELDS},
            "updated_at": now_iso(),
        },
    )
    settings, guides, _ = load_learning_guides()
    guides[guide_key(kind, book, number)] = guide
    save_database_settings(settings)
    return jsonify({
        "guide": guide,
        "guides": public_learning_guides(),
    })


@learning_coaching_bp.put("/guides/<book>/<int:stage>")
@login_required
def update_stage_guide(book, stage):
    if not can_manage_accounts():
        return jsonify({"error": "只有管理员可以维护辅导重点。"}), 403
    if book not in LEARNING_BOOKS or stage < 1 or stage > 3:
        return jsonify({"error": "辅导阶段不存在。"}), 404

    payload = request.get_json(silent=True) or {}
    guide = normalize_stage_guide(
        book,
        stage,
        {
            "focus": str(payload.get("focus") or "").strip()[:5000],
            "practice": str(payload.get("practice") or "").strip()[:5000],
            "updated_at": now_iso(),
        },
    )
    settings, guides, _ = load_stage_guides()
    guides[stage_guide_key(book, stage)] = guide
    save_database_settings(settings)
    return jsonify({
        "guide": guide,
        "stage_guides": public_stage_guides(),
    })


@learning_coaching_bp.post("/<class_id>/upload")
@login_required
def upload_learning_scores(class_id):
    file_storage = request.files.get("file")
    if not file_storage:
        return jsonify({"error": "请上传 Excel 或 CSV 文件。"}), 400

    store = load_store()
    target_class = find_class_by_id(store, class_id)
    if target_class is None or not can_read_class(target_class):
        return jsonify({"error": "班级不存在。"}), 404
    if not can_write_learning_class(target_class):
        return jsonify({"error": "只能上传自己班级的学情数据。"}), 403

    activity_store = load_activity_store(store)
    active_activity = active_completion_activity(activity_store)
    completion_period = completion_period_from_store(store, current_month_key(), active_activity)
    try:
        imported_students = parse_upload(file_storage, completion_period)
    except ValueError as error:
        return jsonify({"error": str(error)}), 400

    updated_at = now_iso()
    month_key = completion_period.get("month") or current_month_key()
    roster = prepare_learning_roster(target_class, imported_students, updated_at)
    if roster.get("error"):
        return jsonify({"error": roster["error"]}), roster.get("status", 400)

    result = merge_learning_upload(
        target_class,
        imported_students,
        month_key,
        updated_at,
        roster_accounts=roster["accounts"],
    )
    if result["score_count"] <= 0:
        return jsonify({"error": "没有识别到可更新的检测分数，请确认表头包含类似 PU1 Unit1 单元检测、PU1 Unit1-3 阶段测评。"}), 400

    removed_appointments = remove_learning_appointments_for_students(class_id, roster["hidden_student_ids"])
    result["roster_initialized"] = roster["is_initial"]
    result["roster_removed_count"] = len(roster["hidden_student_ids"])
    result["removed_appointments"] = removed_appointments

    target_class["updated_at"] = updated_at
    save_store(store)
    return jsonify({
        "result": result,
        "class": class_learning_payload(target_class, public_learning_rounds()),
    })

@learning_coaching_bp.post("/<class_id>/remove-roster-students")
@login_required
def remove_learning_roster_students(class_id):
    payload = request.get_json(silent=True) or {}
    requested_ids = {
        str(student_id or "").strip()
        for student_id in (payload.get("student_ids") or [])
        if str(student_id or "").strip()
    }
    if not requested_ids:
        return jsonify({"error": "请先勾选需要移除的学员。"}), 400

    store = load_store()
    target_class = find_class_by_id(store, class_id)
    if target_class is None or not can_read_class(target_class):
        return jsonify({"error": "班级不存在。"}), 404
    if not can_write_learning_class(target_class):
        return jsonify({"error": "只能维护自己班级的辅导名单。"}), 403

    visible_ids = {
        str(student.get("id") or "").strip()
        for student in learning_roster_students(target_class)
        if str(student.get("id") or "").strip()
    }
    target_ids = requested_ids & visible_ids
    if not target_ids:
        return jsonify({"error": "所选学员已不在当前辅导名单中。"}), 400

    for student in target_class.get("students", []):
        if str(student.get("id") or "").strip() in target_ids:
            student["learning_coaching_hidden"] = True
            student["updated_at"] = now_iso()

    removed_appointments = remove_learning_appointments_for_students(class_id, target_ids)
    target_class["updated_at"] = now_iso()
    save_store(store)
    return jsonify({
        "result": {
            "removed_count": len(target_ids),
            "removed_appointments": removed_appointments,
        },
        "class": class_learning_payload(target_class, public_learning_rounds()),
    })
